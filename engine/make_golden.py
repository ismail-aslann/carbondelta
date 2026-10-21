"""Golden test vektörlerini ve Excel doğrulama dosyasını üretir.

Çalıştır: python -m engine.make_golden
Çıktı: shared/golden/vectors.json (Python ↔ Dart ortak test), excel/dogrulama.xlsx
"""
from __future__ import annotations

import json
import random
from dataclasses import asdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from engine import data
from engine import packages as pk
from engine.core import (NO_LOAD, NON_SHIFTABLE, OK, RESTRICTED_CONFLICT, SAME_WINDOW,
                         simulate)

SEED = 42
ROOT = Path(__file__).resolve().parent.parent
GOLDEN_FILE = ROOT / "shared" / "golden" / "vectors.json"
EXCEL_FILE = ROOT / "excel" / "dogrulama.xlsx"

# (tür, adet) — toplam 30
PLAN = [
    ("non_shiftable", 3),
    ("restricted_conflict", 3),
    ("parametric", 2),
    ("no_load", 2),
    ("same_window", 2),
    ("ok", 18),
]
P_CHOICES = [1.0, 1.0, 0.75, 0.5, 0.25, 0.6]


def _pick_day(rng: random.Random, asset_id: str, month_cycle: list[str]) -> str:
    days = data.days_for(asset_id)
    months = sorted({d[:7] for d in days})
    # Ay dağılımını yaymak için önce henüz kullanılmamış ayları dene.
    fresh = [m for m in months if m not in month_cycle]
    month = rng.choice(fresh or months)
    month_cycle.append(month)
    return rng.choice([d for d in days if d.startswith(month)])


def _scenario(rng, kind, used_months, shiftable_ids, fixed_ids):
    """Tür koşulunu sağlayan bir senaryo bulana kadar dene."""
    for _ in range(10_000):
        ci_kind = rng.choice(["tCO2", "tCO2", "tCO2esd"])
        mode = "parametric" if kind == "parametric" else "data"
        asset_id = rng.choice(fixed_ids if kind == "non_shiftable" else shiftable_ids)
        date = _pick_day(rng, asset_id, used_months)
        machine = data.machine(asset_id)
        ci = data.ci_24(date, ci_kind)
        param_kwh = None
        if mode == "parametric":
            param_kwh = float(rng.choice([25.0, 40.5, 80.0, 112.25]))
            load = [param_kwh] * 24
        else:
            load = data.load_24(asset_id, date)

        L = rng.randint(1, 6)
        s = rng.randint(0, 24 - L)
        t = s if kind == "same_window" else rng.randint(0, 24 - L)
        p = rng.choice(P_CHOICES)
        restricted: list[int] = []
        if kind == "restricted_conflict":
            restricted = sorted({rng.randint(t, t + L - 1), rng.randint(0, 23)})
        elif kind in ("ok", "parametric") and rng.random() < 0.4:
            restricted = sorted(h for h in rng.sample(range(24), 3) if not t <= h < t + L)

        r = simulate(load, ci, s, L, t, p, set(restricted), machine["shiftable"])
        wanted = {
            "non_shiftable": NON_SHIFTABLE,
            "restricted_conflict": RESTRICTED_CONFLICT,
            "parametric": OK,
            "no_load": NO_LOAD,
            "same_window": SAME_WINDOW,
            "ok": OK,
        }[kind]
        if r.status != wanted:
            used_months.pop()
            continue
        if kind == "same_window" and r.baseline_kg == 0.0:
            used_months.pop()
            continue
        return {
            "kind": kind,
            "asset_id": asset_id,
            "date": date,
            "ci_kind": ci_kind,
            "mode": mode,
            "param_kwh": param_kwh,
            "inputs": {
                "load_kwh": load,
                "ci": ci,
                "s": s,
                "L": L,
                "t": t,
                "p": p,
                "restricted": restricted,
                "shiftable": machine["shiftable"],
            },
            "expected": asdict(r),
        }
    raise RuntimeError(f"{kind} için senaryo bulunamadı")


def _require_base_data() -> None:
    """Golden vektörler mobil tarafla paylaşılır: yalnızca varsayılan (yerleşik) veriyle üretilmeli."""
    active = [p.name for p in pk.list_packages() if p.enabled and p.kind != pk.KIND_BROKEN]
    if active:
        raise RuntimeError("Golden vektörler yalnızca varsayılan veriyle üretilir; etkin kullanıcı paketleri var: "
                           + ", ".join(active) + ". 'python -m engine.make_golden' bunları otomatik yok sayar.")


def build_vectors() -> list[dict]:
    _require_base_data()
    rng = random.Random(SEED)
    machines = data.machines()
    shiftable_ids = [m["id"] for m in machines if m["shiftable"]]
    fixed_ids = [m["id"] for m in machines if not m["shiftable"]]
    used_months: list[str] = []
    scenarios = []
    for kind, count in PLAN:
        for _ in range(count):
            sc = _scenario(rng, kind, used_months, shiftable_ids, fixed_ids)
            sc = {"id": f"G{len(scenarios) + 1:02d}", **sc}
            scenarios.append(sc)
    return scenarios


def write_golden(scenarios: list[dict]) -> None:
    GOLDEN_FILE.parent.mkdir(parents=True, exist_ok=True)
    doc = {
        "version": 1,
        "seed": SEED,
        "tolerance": 1e-9,
        "formula": "shared/FORMUL.md v1",
        "scenarios": scenarios,
    }
    # json, float'ları repr() ile yazar → tam double hassasiyeti.
    with open(GOLDEN_FILE, "w", encoding="utf-8", newline="\n") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)


def write_excel(scenarios: list[dict]) -> None:
    EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Dogrulama"

    head = ["Senaryo", "Makine", "Gün", "Mod", "s", "L", "t", "p",
            "B Excel", "S Excel", "Delta % Excel",
            "B Python", "S Python", "Delta % Python",
            "fark B", "fark S", "fark Delta %"]
    kwh_cols = [f"kWh_{h}" for h in range(24)]
    ci_cols = [f"CI_{h}" for h in range(24)]
    ws.append(head + kwh_cols + ci_cols)
    for c in ws[1]:
        c.font = Font(bold=True)

    k0 = len(head) + 1          # ilk kWh sütunu
    c0 = k0 + 24                # ilk CI sütunu
    kL = get_column_letter(k0)
    cL = get_column_letter(c0)

    row_i = 1
    for sc in scenarios:
        if sc["expected"]["status"] != OK:
            continue
        row_i += 1
        inp, exp = sc["inputs"], sc["expected"]
        r = row_i
        kwh_src = f"OFFSET(${kL}{r},0,E{r},1,F{r})"
        ci_src = f"OFFSET(${cL}{r},0,E{r},1,F{r})"
        ci_tgt = f"OFFSET(${cL}{r},0,G{r},1,F{r})"
        row = [
            sc["id"], sc["asset_id"], sc["date"], sc["mode"],
            inp["s"], inp["L"], inp["t"], inp["p"],
            f"=SUMPRODUCT(H{r}*{kwh_src},{ci_src})",
            f"=SUMPRODUCT(H{r}*{kwh_src},{ci_tgt})",
            f"=(J{r}-I{r})/I{r}*100",
            exp["baseline_kg"], exp["scenario_kg"], exp["delta_pct"],
            f"=ROUND(I{r}-L{r},2)", f"=ROUND(J{r}-M{r},2)", f"=ROUND(K{r}-N{r},2)",
        ]
        ws.append(row + list(inp["load_kwh"]) + list(inp["ci"]))
        for col in range(9, 15):
            ws.cell(r, col).number_format = "0.000000"
        for col in range(15, 18):
            ws.cell(r, col).number_format = "0.00"
            ws.cell(r, col).fill = PatternFill("solid", fgColor="E2EFDA")

    for col in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B2"

    note = wb.create_sheet("Aciklama")
    note.append(["CarbonDelta — Excel doğrulaması (shared/FORMUL.md v1)"])
    note.append(["B = SUMPRODUCT(p × kWh[s..s+L-1], CI[s..s+L-1]); S = aynı yük × CI[t..t+L-1]"])
    note.append(["fark sütunları 2 ondalıkta 0,00 olmalıdır."])
    note.append(["Örnek yük, Balıkesir tesisi değil · Kaynak: EPİAŞ 2025, ETKB 2023, Zenodo"])
    wb.save(EXCEL_FILE)


def main() -> None:
    # Kullanıcının %APPDATA% veri paketleri golden'a karışmasın: boş geçici veri klasörü.
    import os
    import tempfile
    os.environ["CARBONDELTA_DATA_DIR"] = tempfile.mkdtemp(prefix="carbondelta_golden_")
    data.reload()
    scenarios = build_vectors()
    write_golden(scenarios)
    write_excel(scenarios)
    counts: dict[str, int] = {}
    for sc in scenarios:
        counts[sc["expected"]["status"]] = counts.get(sc["expected"]["status"], 0) + 1
    print(f"Senaryo: {len(scenarios)} {counts}")
    print(f"Aylar: {sorted({sc['date'][:7] for sc in scenarios})}")
    print(f"Yazıldı: {GOLDEN_FILE.relative_to(ROOT)}, {EXCEL_FILE.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
