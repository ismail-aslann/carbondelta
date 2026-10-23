"""Proje Pazarı tam metin taslağını üretir: basvuru/CarbonDelta_Tam_Metin_Taslak.docx

Rakamlar sonuclar.json'dan gelir (önce sonuclari_hesapla.py çalıştırılır).
Çalıştırma: basvuru\\.araclar\\Scripts\\python basvuru\\uret_tam_metin.py
[[...]] ile yazılanlar sarı vurgulu yer tutuculardır; gönderimden önce doldurulmalıdır.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
R = json.loads((HERE / "sonuclar.json").read_text(encoding="utf-8"))
IMG = HERE / "gorseller"
OUT = HERE / "CarbonDelta_Tam_Metin_Taslak.docx"
FONT = "Times New Roman"


def num(v, d=2):
    s = f"{abs(v):,.{d}f}".replace(",", " ").replace(".", ",").replace(" ", ".")
    return ("−" if round(v, d) < 0 else "") + s


def pct(v, d=2):
    return f"%{num(v, d)}" if v >= 0 else f"−%{num(-v, d)}"


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 12), ("Heading 2", 11)):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)


def runs(par, text, size=None, italic=False):
    """**kalın** ve [[yer tutucu]] işaretlerini çözer."""
    for part in re.split(r"(\*\*.+?\*\*|\[\[.+?\]\])", text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("[["):
            r = par.add_run(part[2:-2])
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = par.add_run(part)
        r.italic = italic or r.italic
        if size:
            r.font.size = Pt(size)
    return par


def p(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, **kw):
    par = doc.add_paragraph()
    par.alignment = align
    return runs(par, text, **kw)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def bullets(items):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        runs(par, it)


def formula(text):
    p(text, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)


FIG = [0]
TAB = [0]


def figure(name, caption, width=15.5):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(IMG / name), width=Cm(width))
    FIG[0] += 1
    p(f"**Şekil {FIG[0]}.** {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(headers, rows, caption, widths):
    TAB[0] += 1
    p(f"**Tablo {TAB[0]}.** {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        runs(c.paragraphs[0], f"**{htext}**", size=9.5)
        shade(c, "E7F3EF")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            runs(cells[i].paragraphs[0], str(v), size=9.5)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        for i, c in enumerate(row.cells):
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


# ------------------------------------------------------------------ içerik
prof, dh, ch = R["profile"], R["dirtiest_h"], R["cleanest_h"]
spread = (prof[dh] - prof[ch]) / prof[dh] * 100
ref = R["ref_day"]
comp = next(m for m in R["machines"] if m["id"] == "comp_b")
tot = R["total"]
avoided_per_mwh = (prof[dh] - prof[ch]) * 1000  # kg CO2 / MWh

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("CarbonDelta: Endüstriyel Elektrik Yüklerini Daha Düşük Karbonlu Saatlere Kaydırmak İçin "
                  "Yapay Zekâ Destekli Karar Destek Yazılımı")
r.bold, r.font.size = True, Pt(14)
p("[[Sunumu yapacak kişi (ilk isim)]], İsmail Ali ASLAN, Emin Burak ÖZTÜRK, Öğr. Gör. Ali ERFİDAN (Danışman)",
  align=WD_ALIGN_PARAGRAPH.CENTER)
p("Balıkesir Üniversitesi, [[Fakülte / Bölüm]] · [[iletişim e-postası]]", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
p("**Tematik alan:** Yeşil Teknolojiler ve Döngüsel Ekonomi (karbon yönetimi)", align=WD_ALIGN_PARAGRAPH.CENTER,
  size=10)

h1("Özet")
p("Sanayide elektrik tüketiminden kaynaklanan dolaylı emisyonlar, yalnızca ne kadar elektrik kullanıldığına değil, "
  "elektriğin hangi saatte kullanıldığına da bağlıdır; çünkü şebekedeki üretim karışımı gün içinde sürekli değişir. "
  "CarbonDelta, kompresör, pres ve karıştırıcı gibi zamanı esnetilebilen yüklerin gün içinde başka bir zaman "
  "penceresine alınmasının emisyona etkisini saatlik çözünürlükte hesaplayan, internet bağlantısı gerektirmeyen bir "
  "masaüstü karar destek yazılımıdır. Yazılım, EPİAŞ Şeffaflık Platformu'nun 2025 yılına ait saatlik gerçek zamanlı "
  "üretim verisini Enerji ve Tabii Kaynaklar Bakanlığı'nın (ETKB) 2023 emisyon faktörleriyle birleştirerek her saat "
  "için şebeke karbon yoğunluğu üretir ve makine bazlı yük verisiyle standart ve hedef penceredeki emisyonları "
  "karşılaştırır. Hesap çekirdeği belirlenimcidir ve Excel doğrulama dosyası ile otomatik test vektörleriyle aynı "
  "sonucu verir. Uygulama günün en temiz penceresini önerir, aynı kaydırmanın bir yıl boyunca uygulanması "
  "durumundaki toplam etkiyi aylık olarak hesaplar ve işletmenin kendi ölçümlerini yazılımı değiştirmeden içe "
  "aktarmasına izin verir. “Delti” adlı yapay zekâ yardımcısı yalnızca hesaplanmış sayıları açıklama "
  "metnine dönüştürür; metindeki her sayı hesap sonucuyla otomatik olarak karşılaştırılır. 2025 verisinde yıllık "
  f"ortalama karbon yoğunluğu {dh:02d}:00'da {num(prof[dh], 3)}, {ch:02d}:00'da {num(prof[ch], 3)} tCO2/MWh olup "
  f"aradaki fark yaklaşık %{num(spread, 1)}'dir. Örnek kompresör yükünün 03:00–06:00 aralığından 18:00–21:00 "
  f"aralığına alınması, kaydırılan enerjinin emisyonunu örnek günde {pct(ref['delta_pct'])}, yıl genelinde "
  f"{pct(comp['delta_pct'])} değiştirmektedir. CarbonDelta, yatırım gerektirmeden yalnızca zamanlama kararıyla elde "
  "edilebilecek emisyon azaltımını görünür kılarak sanayinin yeşil ve dijital dönüşümüne katkı sunmayı amaçlar.")
p("**Anahtar kelimeler:** karbon yoğunluğu, yük kaydırma, talep tarafı esnekliği, karar destek sistemi, "
  "Kapsam 2 emisyonları, ikiz dönüşüm")

h1("1. Giriş ve Problem Tanımı")
p("Elektrik şebekesinde talep, her an farklı kaynaklardan karşılanır. Gece saatlerinde baz yükü taşıyan kömür "
  "santrallerinin payı artarken, gün içinde ve akşam saatlerinde yenilenebilir kaynakların ve doğal gazın payı "
  "değişir. Bu nedenle aynı miktarda elektriğin şebekeye yüklediği karbon, tüketildiği saate göre farklıdır. "
  "Sanayi tesislerinde ise basınçlı hava üretimi, ön ısıtma, karıştırma, şarj ve bazı pres işleri gibi yüklerin "
  "zamanlaması çoğu zaman üretim kalitesini etkilemeden birkaç saat öne ya da arkaya alınabilir.")
p("Bu esneklik, ek bir yatırım gerektirmeyen bir emisyon azaltım imkânı sunar; ancak tesis yöneticileri hangi yükü "
  "hangi saate taşıdıklarında ne kadar etki elde edeceklerini hızlı, şeffaf ve doğrulanabilir biçimde hesaplayacak "
  "araçlara çoğu zaman sahip değildir. Kurumsal karbon muhasebesinde genellikle yıllık ortalama emisyon faktörü "
  "kullanıldığından, saatten saate oluşan fark hesaplarda görünmez.")
p("Konu düzenleyici açıdan da önem kazanmaktadır. Temmuz 2025'te yürürlüğe giren 7552 sayılı İklim Kanunu, "
  "Türkiye Emisyon Ticaret Sistemi'nin kurulmasının dayanağını oluşturmuş ve sistem pilot uygulama dönemiyle "
  "başlamıştır. Avrupa Birliği'nin Sınırda Karbon Düzenleme Mekanizması (SKDM/CBAM) 1 Ocak 2026 itibarıyla mali "
  "yükümlülük içeren dönemine geçmiştir; çimento ve gübre gibi ürünlerde elektrik tüketiminden kaynaklanan dolaylı "
  "emisyonlar da hesaba katılmaktadır. İhracat yapan sanayi kuruluşları için ürün başına emisyonu ölçmek ve azaltmak "
  "giderek bir rekabet koşuluna dönüşmektedir.")

h1("2. Amaç ve Kapsam")
p("Projenin amacı, sanayi kullanıcısının bir makinenin yükünü gün içinde başka bir zaman penceresine aldığında "
  "şebeke kaynaklı emisyonun nasıl değişeceğini saatlik verilerle hesaplayan, sonuçları sade bir dille açıklayan "
  "ve kararı kullanıcıya bırakan bir karar destek yazılımı geliştirmektir. Yazılımın temel hedefleri şunlardır:")
bullets([
    "Kaydırma senaryosunun emisyon etkisini resmî veriye dayalı, tekrarlanabilir ve denetlenebilir biçimde hesaplamak,",
    "Günün ve yılın en temiz zaman penceresini otomatik olarak önermek,",
    "Aynı kaydırmanın bir yıl boyunca uygulanması durumundaki toplam etkiyi aylık kırılımla göstermek,",
    "İşletmenin kendi makine ölçümlerini, yeni yıl üretim verisini ve güncel emisyon faktörlerini kod değişikliği "
    "olmadan eklemesine imkân vermek,",
    "Sonuçları yapay zekâ ile anlaşılır bir metne dönüştürürken sayısal doğruluğu güvence altına almak.",
])
p("Yazılım karar destek amaçlıdır; hiçbir makineye doğrudan müdahale etmez. Parasal kazanç hesabı ve SKDM "
  "beyan hesabı bu sürümün kapsamı dışındadır.")

h1("3. Yöntem")
h2("3.1. Veri kaynakları")
bullets([
    f"**EPİAŞ Şeffaflık Platformu, 2025:** {R['grid_days']} günlük (8.760 saat) kaynak bazında gerçek zamanlı "
    "üretim verisi (doğal gaz, linyit, ithal kömür, rüzgâr, güneş, hidro vb.).",
    "**ETKB 2023 emisyon faktörleri:** 15 yakıt türü için resmî tCO2/MWh ve tCO2-eşd/MWh değerleri.",
    "**Makine yükleri:** Bir imalat tesisine ait 15 dakikalık makine bazlı enerji ölçümlerinden oluşan açık veri "
    "seti (Zenodo); 7 makine, temizlenmiş ve saatlik toplama dönüştürülmüş veri. Bu veri yöntemin gösterimi için "
    "örnek yük olarak kullanılmıştır; Balıkesir'deki bir tesisin ölçümü değildir.",
])
h2("3.2. Saatlik karbon yoğunluğu")
p("Her saatin şebeke karbon yoğunluğu, o saatteki kaynak üretimlerinin ilgili emisyon faktörüyle ağırlıklandırılıp "
  "toplam yurt içi üretime bölünmesiyle hesaplanır:")
formula("CI(h) = Σk [ Gk(h) × EFk ] / Σk Gk(h)")
p("Burada Gk(h), k kaynağının h saatindeki üretimini (MWh), EFk ise k kaynağının emisyon faktörünü (tCO2/MWh) "
  "gösterir. Yenilenebilir kaynakların doğrudan yakıt emisyonu sıfır kabul edilir.")
h2("3.3. Kaydırma modeli")
p("Kullanıcı bir makine, gün, standart başlangıç saati s, pencere uzunluğu L, hedef başlangıç saati t ve kaydırma "
  "oranı p (0–1) seçer. Standart penceredeki her saatin enerjisinin p kadarı hedef pencerede aynı sırayla tüketilir:")
formula("Ei = p × kWh(s+i),   i = 0 … L−1")
formula("B = Σ Ei × CI(s+i)        S = Σ Ei × CI(t+i)")
formula("Değişim (%) = (S − B) / B × 100        Günlük etki (%) = (S − B) / Günlük toplam emisyon × 100")
p("Kaydırılamayan süreçler (havalandırma, pres ısıtma), gün dışına taşan pencereler, kullanıcının kısıtladığı "
  "saatlerle çakışan hedefler ve standart pencerede yük bulunmayan günler ayrı durum kodlarıyla raporlanır. "
  "Yuvarlama yalnızca gösterimde yapılır; tüm hesaplar tam hassasiyetle yürütülür.")
h2("3.4. En temiz pencere ve yıllık etki")
p("En temiz pencere, kısıtlı saatler dışında kalan tüm geçerli hedef başlangıçları taranarak S değeri en düşük "
  "olan pencere olarak belirlenir. Yıllık etki için aynı senaryo, seçilen yılda verisi olan her günde aynı formülle "
  "hesaplanır ve toplanır:")
formula("Yıllık değişim (%) = (ΣS − ΣB) / ΣB × 100")
p("Yazılım ayrıca yıl boyunca sabit tutulduğunda toplam emisyonu en düşük olan hedef saati ve aylık kırılımı "
  "gösterir. Verisi olmayan günler tahminle tamamlanmaz.")
h2("3.5. Yazılım mimarisi")
p("CarbonDelta Python ile geliştirilmiş, PySide6 (Qt) arayüzlü bir Windows masaüstü uygulamasıdır. Hesap çekirdeği "
  "arayüzden bağımsız bir modüldür; aynı formül bir Excel doğrulama dosyası ve sabit test vektörleriyle (referans "
  "senaryo: Kompresör, 15.07.2025, 03:00→18:00) karşılaştırılır. Projede 138 otomatik test bulunmaktadır. Uygulama "
  "tek dosyalık bir kurulum sihirbazıyla dağıtılır, yönetici yetkisi gerektirmez ve hesap için internet bağlantısına "
  "ihtiyaç duymaz. Kaydedilen senaryolar yerel bir SQLite veritabanında tutulur. Arayüz Türkçedir ve fare "
  "kullanılmadan tamamen klavyeyle yönetilebilir.")
h2("3.6. Veri ekleme")
p("Kullanıcı, kendi tesisine ait saatlik veya 15 dakikalık makine ölçümlerini (CSV/Excel), EPİAŞ'tan indirdiği yeni "
  "yıl üretim verisini ve yeni emisyon faktörlerini “Veri” ekranından ekleyebilir. Uygulama her günün 24 "
  "tekrarsız saatini, eksik ve negatif değerleri ve şebeke verisiyle eşleşmeyi denetler; sorunlu günleri raporlayarak "
  "dışarıda bırakır. Gerçek tesis ölçümüyle hesaplanan sonuçlar ekranda “Tesis ölçümü” etiketiyle, örnek "
  "veriyle hesaplananlar “Örnek yük” etiketiyle ayrıştırılır.")
h2("3.7. Yapay zekâ katmanı ve sayısal doğrulama")
p("“Delti” adlı yardımcı, bir büyük dil modeli aracılığıyla iki görev üstlenir: hesap sonuçlarını sade "
  "Türkçe ile açıklamak ve “Kompresörü 3'ten 18'e al” gibi günlük dildeki cümleleri senaryo girdilerine "
  "dönüştürmek. Modele yalnızca hesap çekirdeğinin ürettiği sayılar gönderilir; model hesap yapmaz. Üretilen "
  "metindeki her sayı, girdilerdeki değerlerle otomatik olarak karşılaştırılır; veride bulunmayan bir sayı tespit "
  "edilirse yapay zekâ metni gösterilmez, yerine şablon metin kullanılır. İnternet bağlantısı veya API anahtarı "
  "yoksa uygulama tüm hesap işlevleriyle çalışmaya devam eder.")

h1("4. Bulgular")
h2("4.1. Saatlik karbon yoğunluğu profili")
p(f"2025 verisinin yıllık ortalamasında en yüksek karbon yoğunluğu {dh:02d}:00 saatinde ({num(prof[dh], 3)} "
  f"tCO2/MWh), en düşük değer {ch:02d}:00 saatinde ({num(prof[ch], 3)} tCO2/MWh) görülmektedir (Şekil 1). Gece ve "
  f"sabahın erken saatleri ile akşam saatleri arasındaki yaklaşık %{num(spread, 1)}'lik fark, yük kaydırmanın "
  "potansiyelini ortaya koymaktadır.")
figure("07_saatlik_profil.png", "2025 yılı yıllık ortalama saatlik şebeke karbon yoğunluğu.")
h2("4.2. Örnek gün senaryosu")
p("Referans senaryoda Kompresör yükünün tamamı 15.07.2025 tarihinde 03:00–06:00 aralığından 18:00–21:00 aralığına "
  f"alınmıştır. Kaydırılan {num(ref['kwh'])} kWh enerjinin emisyonu {num(ref['B'])} kgCO2'den {num(ref['S'])} "
  f"kgCO2'ye düşmekte, değişim {pct(ref['delta_pct'])} olmaktadır. Bu, makinenin o günkü toplam emisyonunda "
  f"{pct(ref['daily_pct'])}'lik bir azalmaya karşılık gelir (Şekil 2).")
figure("02_senaryo.png", "Senaryo ekranı: makine yükü, saatlik karbon yoğunluğu ve sonuç.")
h2("4.3. Yıllık etki")
p("Aynı kaydırma (03:00–06:00 → 18:00–21:00, %100) 2025 yılında verisi bulunan tüm günlerde uygulandığında elde "
  "edilen sonuçlar Tablo 1'de verilmiştir. Kaydırılabilir beş makinenin tamamında emisyon azalmakta ve yıl boyunca "
  "sabit tutulabilecek en temiz hedef pencere 18:00–21:00 olarak bulunmaktadır.")
rows = []
for m in R["machines"]:
    if m["shiftable"]:
        rows.append([m["name"], f"{m['days_ok']} / {m['days']}", num(m["shifted_kwh"], 0), num(m["delta_kg"], 1),
                     pct(m["delta_pct"]), pct(m["daily_pct"])])
    else:
        rows.append([m["name"], f"— / {m['days']}", "kaydırılamaz", "—", "—", "—"])
rows.append(["**Kaydırılabilir makineler toplamı**", "", "", f"**{num(tot['delta_kg'], 1)}**",
             f"**{pct(tot['delta_pct'])}**", f"**{pct(tot['daily_pct'])}**"])
table(["Makine", "Hesaba giren gün", "Kaydırılan enerji (kWh/yıl)", "Değişim (kg CO2/yıl)",
       "Kaydırılan yükte değişim", "Toplam emisyonda pay"],
      rows, "Örnek yüklerle 2025 yıllık etki (03:00–06:00 → 18:00–21:00).", [4.2, 2.2, 2.6, 2.4, 2.3, 2.3])
better_txt = ("tamamında emisyon azalmış, hiçbir günde artış görülmemiştir" if comp["worse"] == 0
              else f"{comp['better']} gününde emisyon azalmış, {comp['worse']} gününde artmıştır")
p(f"Kompresör için hesaba giren {comp['days_ok']} günün {better_txt}. Aylık kırılım, etkinin yaz ve sonbahar "
  "aylarında belirginleştiğini göstermektedir (Şekil 3).")
figure("08_aylik_etki.png", "Kompresör için 2025 aylık emisyon değişimi.")
h2("4.4. Ölçeklenebilirlik için basit örnek")
p("Örnek veri setindeki makineler küçük güçlü olduğundan mutlak değerler düşüktür. Yıllık ortalama profile göre "
  f"{dh:02d}:00 saatinden {ch:02d}:00 saatine kaydırılan her 1 MWh enerji yaklaşık {num(avoided_per_mwh, 0)} kg "
  f"CO2 daha az emisyona karşılık gelir. Örneğin günde 1 MWh esnek yükü olan bir tesis, bu kaydırmayı yılda 300 "
  f"iş gününde uyguladığında yaklaşık {num(avoided_per_mwh * 300 / 1000, 1)} ton CO2 azaltım potansiyeline ulaşır. "
  "Bu hesap ortalama yoğunluklara dayalı bir göstergedir; tesise özgü sonuç, CarbonDelta'ya tesisin kendi ölçümleri "
  "eklenerek elde edilir.")
figure("04_yillik_etki.png", "Yıllık etki ekranı: özet göstergeler, aylık tablo ve grafik.")

h1("5. Yenilikçi Yönler ve Özgün Değer")
bullets([
    "**Saatlik ve makine bazlı hesap:** Yıllık ortalama faktör yerine her saatin gerçek üretim karışımı kullanılır; "
    "karar doğrudan bir makinenin bir zaman penceresi düzeyinde verilir.",
    "**Doğrulanmış yapay zekâ:** Dil modeli hesap yapmaz; ürettiği her sayı otomatik denetlenir. Böylece yapay "
    "zekânın kullanıcı dostu anlatımı ile hesabın güvenilirliği bir arada sunulur.",
    "**Tekrarlanabilirlik:** Kilitli formül, Excel doğrulama dosyası ve test vektörleri sayesinde aynı girdiler her "
    "ortamda aynı sonucu verir; sonuçlar denetime açıktır.",
    "**Çevrim dışı çalışma ve kolay kurulum:** Hesap için internet gerekmez; tek dosyayla kurulur.",
    "**Kendi verisiyle büyüyen yapı:** Yeni makine, yeni yıl ve yeni emisyon faktörleri kullanıcı tarafından "
    "eklenebilir; yazılım her yıl yeniden geliştirme gerektirmez.",
    "**Yıllık planlama:** Tek günlük sonucun yanında, bir kaydırma alışkanlığının yıl boyunca sağlayacağı toplam "
    "etki ve en uygun sabit hedef saat gösterilir.",
])

h1("6. Ticarileşme Potansiyeli ve Hedef Kitle")
p("Birincil hedef kitle; enerji yoğun KOBİ'ler, organize sanayi bölgelerindeki üreticiler, AB'ye ihracat yapan ve "
  "ürün karbon ayak izini raporlaması gereken firmalar ile enerji yönetimi ve sürdürülebilirlik danışmanlığı "
  "veren kuruluşlardır. Enerji yönetim sistemi (ISO 50001) ve sürdürülebilirlik raporu hazırlayan işletmeler, "
  "CarbonDelta çıktılarını azaltım eylemlerini önceliklendirmek için kullanabilir.")
p("Öngörülen iş modeli; yıllık yazılım lisansı, tesis verisinin aktarımı ve ilk analiz için kurulum hizmeti ile "
  "danışmanlık firmalarına yönelik çoklu tesis lisansından oluşmaktadır. İlk aşamada Balıkesir'deki sanayi "
  "kuruluşlarıyla, üniversitenin teknoloji transfer ofisi aracılığıyla pilot uygulamalar yapılması "
  "hedeflenmektedir. Genel amaçlı karbon muhasebesi yazılımları yıllık toplam emisyonu raporlamaya odaklanırken, "
  "CarbonDelta zamanlama kararına odaklanır ve bu araçların tamamlayıcısı olarak konumlanır.")

h1("7. İkiz Dönüşüm ve Sürdürülebilirlik Katkısı")
p("Proje, etkinliğin teması olan ikiz dönüşümün iki bileşenini birlikte ele alır: veri analitiği ve yapay zekâ "
  "(dijital dönüşüm) kullanılarak elektrik tüketiminden kaynaklanan emisyonların azaltılması (yeşil dönüşüm). "
  "Yazılım, yatırım ya da ekipman değişikliği gerektirmeden mevcut tüketimin zamanlamasını iyileştirerek "
  "yenilenebilir kaynakların daha yüksek paya sahip olduğu saatlerin değerlendirilmesine katkı sağlar. Bu yönüyle "
  "Birleşmiş Milletler Sürdürülebilir Kalkınma Amaçları'ndan Erişilebilir ve Temiz Enerji (7), Sanayi, Yenilikçilik "
  "ve Altyapı (9), Sorumlu Üretim ve Tüketim (12) ve İklim Eylemi (13) ile doğrudan ilişkilidir.")

h1("8. Sınırlılıklar ve Gelecek Çalışmalar")
bullets([
    "**Ortalama ve marjinal emisyon:** Hesap, saatlik ortalama karbon yoğunluğuyla (konum bazlı yöntem) yapılır. Bir "
    "yükü kaydırmanın şebekedeki gerçek etkisi, o saatte üretimini değiştiren marjinal santrale bağlıdır. Bu nedenle "
    "sonuçlar karşılaştırma ve önceliklendirme için bir gösterge niteliğindedir; marjinal emisyon faktörü desteği "
    "planlanmaktadır.",
    "**Örnek yük:** Kullanılan makine verisi açık bir veri setinden alınmıştır. Bir sonraki adım, Balıkesir'deki bir "
    "tesisin ölçümleriyle saha doğrulamasıdır.",
    "**Faktör ve kapsam:** ETKB 2023 faktörleri 2025 üretimine uygulanmaktadır; elektrik ithalatı/ihracatı ve gerçek "
    "zamanlı üretim verisinde yer almayan üretim hesaba girmez.",
    "**Kaydırma varsayımı:** Kaydırılan enerji miktarının aynı kaldığı ve yükün aynı gün içinde taşındığı kabul "
    "edilir; işletme kısıtları kullanıcı tarafından kısıtlı saat olarak girilir.",
    "**Gelecek çalışmalar:** EPİAŞ verisinin otomatik güncellenmesi, piyasa takas fiyatıyla maliyet etkisinin "
    "gösterilmesi, birden fazla makinenin birlikte planlanması, mobil uygulama ve tesis enerji izleme sistemleriyle "
    "entegrasyon.",
])

h1("9. Sonuç")
p("CarbonDelta, şebekenin saatlik karbon yoğunluğundaki farkı sanayi kullanıcısının anlayabileceği ve "
  f"uygulayabileceği bir karara dönüştürmektedir. 2025 verisinde en kirli ve en temiz saat arasında yaklaşık "
  f"%{num(spread, 1)}'lik fark bulunmakta; örnek yüklerde basit bir zamanlama değişikliği kaydırılan enerjinin "
  f"emisyonunu yıl genelinde yaklaşık {pct(tot['delta_pct'], 1)} değiştirmektedir. Doğrulanmış hesap çekirdeği, "
  "sayısal denetimli yapay zekâ anlatımı ve kullanıcının kendi verisini ekleyebildiği yapısıyla CarbonDelta, "
  "sanayide yatırım gerektirmeyen emisyon azaltım fırsatlarını görünür kılan, ölçeklenebilir bir araç sunmaktadır.")

h1("Kaynakça")
refs = [
    "EPİAŞ (2026). Şeffaflık Platformu – Gerçek Zamanlı Üretim. https://seffaflik.epias.com.tr (Erişim: [[gg.aa.2026]]).",
    "T.C. Enerji ve Tabii Kaynaklar Bakanlığı. 2023 Yılı Türkiye Elektrik Üretimi ve Elektrik Tüketim Noktası "
    "Emisyon Faktörleri. Ankara.",
    "Flynn, C., Murphy, T., Walsh, J. Industrial Asset-Level Electrical Energy Dataset from a Manufacturing Facility "
    "(15-Minute Aggregates). Zenodo. [[DOI eklenecek]]",
    "7552 sayılı İklim Kanunu. Resmî Gazete, Temmuz 2025.",
    "Regulation (EU) 2023/956 of the European Parliament and of the Council establishing a carbon border adjustment "
    "mechanism. Official Journal of the European Union, L 130, 2023.",
    "World Resources Institute, WBCSD (2015). GHG Protocol Scope 2 Guidance. Washington, DC.",
]
for i, ref_text in enumerate(refs, 1):
    par = p(f"[{i}] {ref_text}", align=WD_ALIGN_PARAGRAPH.LEFT, size=10)
    par.paragraph_format.left_indent = Cm(0.8)
    par.paragraph_format.first_line_indent = Cm(-0.8)

doc.save(OUT)
print("kaydedildi:", OUT)
